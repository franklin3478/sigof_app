import streamlit as st
import requests
from datetime import date
import pandas as pd
from openpyxl import Workbook
from io import BytesIO
import hashlib


# ============================================================
# CONFIGURACIÓN DE API
# ============================================================

URL_BASE = (
    "https://servicios.distriluz.com.pe:51000/"
    "OptimusNGC_FieldService/api"
)

URL_LOGIN = f"{URL_BASE}/auth/login"

URL_CONSULTAR = (
    f"{URL_BASE}/monitoreo/lecturas/consultar"
)

URL_MAPA_RUTA = (
    f"{URL_BASE}/monitoreo/lecturas/mapa/ruta"
)


# ============================================================
# OBSERVACIONES
# ============================================================

OBSERVACIONES_LISTA = [
    "CORRECTO",
    "Sin Epps",
    "Sin Fotocheck",
    "Sin Camisa",
    "Sin Chaleco",
    "Sin Camisa - Sin Chaleco - Sin Fotocheck.",
    "Sin Chaleco - Sin Fotocheck",
    "No se visualiza lecturador",
    "Selfie No corresponde al lecturador",
    "Fotografia incorrecta",
    "Sin Camisa - Sin Gorro",
    "Sin Chaleco - Sin Gorro",
    "Sin Fotocheck - Sin Camisa",
    "Sin Gorro",
    "Relectura"
]


# ============================================================
# CLAVE ÚNICA PARA LA OBSERVACIÓN DE CADA SELFIE
# ============================================================

def clave_observacion(selfie):

    texto = "|".join([
        str(selfie.get("lecturista", "")),
        str(selfie.get("fechaSelfie", "")),
        str(selfie.get("secuenciaEnRuta", "")),
        str(selfie.get("totalRuta", "")),
        str(selfie.get("selfieUrl", ""))
    ])

    return (
        "fs_obs_"
        + hashlib.md5(
            texto.encode("utf-8")
        ).hexdigest()
    )


# ============================================================
# LOGIN FIELDSERVICE
# ============================================================

def login_fieldservice(usuario, clave):

    payload = {
        "usuario": usuario,
        "clave": clave,
        "tipo": 3
    }

    try:

        respuesta = requests.post(
            URL_LOGIN,
            json=payload,
            timeout=30
        )

        if respuesta.status_code != 200:
            return None

        datos = respuesta.json()

        if datos.get("success") is True:

            token = datos.get("accessToken")

            if token:
                return token

        return None

    except requests.exceptions.RequestException:
        return None

    except ValueError:
        return None


# ============================================================
# OBTENER LECTURISTAS
# ============================================================

def obtener_lecturistas(unidad_id, periodo):

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:
        return []

    if unidad_id is None or periodo is None:
        return []

    proveedor_id = 3

    url = (
        f"{URL_BASE}/listageneral/lecturista"
        f"?proveedorId={proveedor_id}"
        f"&unidadId={unidad_id}"
        f"&periodo={periodo}"
    )

    headers = {
        "Authorization": f"Bearer {token}",
        "x-audit-application": "WEB",
        "x-audit-username": "72690389"
    }

    try:

        respuesta = requests.get(
            url,
            headers=headers,
            timeout=30
        )

        if respuesta.status_code != 200:
            return []

        datos = respuesta.json()

        if isinstance(datos, list):
            return datos

        return []

    except requests.exceptions.RequestException:
        return []

    except ValueError:
        return []


# ============================================================
# OBTENER CICLOS
# ============================================================

def obtener_ciclos(unidad_id):

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:
        return []

    if unidad_id is None:
        return []

    empresa_id = 4
    usuario_id = 315
    estado_id = 1

    url = (
        f"{URL_BASE}/listageneral/"
        f"servicio-flujo-ciclo-usuario"
        f"?empresaId={empresa_id}"
        f"&usuarioId={usuario_id}"
        f"&unidadId={unidad_id}"
        f"&estadoId={estado_id}"
    )

    headers = {
        "Authorization": f"Bearer {token}"
    }

    try:

        respuesta = requests.get(
            url,
            headers=headers,
            timeout=30
        )

        if respuesta.status_code != 200:
            return []

        datos = respuesta.json()

        if isinstance(datos, list):
            return datos

        return []

    except requests.exceptions.RequestException:
        return []

    except ValueError:
        return []


# ============================================================
# OBTENER SECTORES
# ============================================================

def obtener_sectores(unidad_id, ciclo_id):

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:
        return []

    if unidad_id is None or ciclo_id is None:
        return []

    empresa_id = 4
    usuario_id = 315

    url = (
        f"{URL_BASE}/listageneral/sector-usuario"
        f"?empresaId={empresa_id}"
        f"&usuarioId={usuario_id}"
        f"&unidadId={unidad_id}"
        f"&cicloId={ciclo_id}"
    )

    headers = {
        "Authorization": f"Bearer {token}"
    }

    try:

        respuesta = requests.get(
            url,
            headers=headers,
            timeout=30
        )

        if respuesta.status_code != 200:
            return []

        datos = respuesta.json()

        if isinstance(datos, list):
            return datos

        return []

    except requests.exceptions.RequestException:
        return []

    except ValueError:
        return []


# ============================================================
# OBTENER RUTAS
# ============================================================

def obtener_rutas(unidad_id, sector_id):

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:
        return []

    if unidad_id is None or sector_id is None:
        return []

    empresa_id = 4
    usuario_id = 315
    estado_id = 1

    url = (
        f"{URL_BASE}/listageneral/ruta-lectura"
        f"?empresaId={empresa_id}"
        f"&sectorId={sector_id}"
        f"&estadoId={estado_id}"
        f"&usuarioId={usuario_id}"
        f"&unidadId={unidad_id}"
    )

    headers = {
        "Authorization": f"Bearer {token}"
    }

    try:

        respuesta = requests.get(
            url,
            headers=headers,
            timeout=30
        )

        if respuesta.status_code != 200:
            return []

        datos = respuesta.json()

        if isinstance(datos, list):
            return datos

        return []

    except requests.exceptions.RequestException:
        return []

    except ValueError:
        return []


# ============================================================
# OBTENER VALOR DE UN CAMPO
# ============================================================

def obtener_valor(item, posibles):

    for campo in posibles:

        if campo in item:

            valor = item.get(campo)

            if valor is not None:
                return valor

    return ""


# ============================================================
# CONVERTIR FECHA A FORMATO API
# ============================================================

def fecha_para_api(valor):

    if valor is None:
        return None

    if hasattr(valor, "strftime"):

        return valor.strftime(
            "%Y-%m-%dT00:00:00"
        )

    texto = str(valor)

    if "T" in texto:

        return (
            texto[:10]
            + "T00:00:00"
        )

    if " " in texto:

        return (
            texto.split(" ")[0]
            + "T00:00:00"
        )

    return (
        texto[:10]
        + "T00:00:00"
    )


# ============================================================
# OBTENER SELFIES DE UNA ASIGNACIÓN
# ============================================================

def obtener_selfies_asignacion(
    token,
    periodo,
    unidad_id,
    fila,
    ciclo_id,
    sector_id,
    ruta_id
):

    proveedor_trabajador_id = obtener_valor(
        fila,
        [
            "proveedorTrabajadorId",
            "proveedorTrabajadorID",
            "trabajadorId",
            "lecturistaId",
            "id"
        ]
    )

    asignacion_id = obtener_valor(
        fila,
        [
            "ordenTrabajoAsignacionId",
            "ordenTrabajoAsignacionID",
            "ordenTrabajoAsignacionIds"
        ]
    )

    fecha_lectura = obtener_valor(
        fila,
        [
            "fechaLectura",
            "fecha",
            "fechaAsignacion"
        ]
    )

    if proveedor_trabajador_id in (
        None,
        ""
    ):
        return []

    if asignacion_id in (
        None,
        ""
    ):
        return []

    if fecha_lectura in (
        None,
        ""
    ):
        return []

    fecha_api = fecha_para_api(
        fecha_lectura
    )

    if not fecha_api:
        return []

    fila_ruta_id = obtener_valor(
        fila,
        [
            "rutaLecturaId"
        ]
    )

    fila_sector_id = obtener_valor(
        fila,
        [
            "sectorId"
        ]
    )

    if fila_ruta_id in (
        None,
        ""
    ):

        fila_ruta_id = ruta_id

    if fila_sector_id in (
        None,
        ""
    ):

        fila_sector_id = sector_id

    payload = {

        "cicloId":
            ciclo_id,

        "fechaAsignacion":
            None,

        "fechaLecturaDesde":
            None,

        "fechaLecturaHasta":
            None,

        "nroServicioDigitado":
            None,

        "ordenTrabajoAsignacionIds":
            str(asignacion_id),

        "ordenTrabajoId":
            None,

        "periodo":
            periodo,

        "proveedorId":
            3,

        "proveedorTrabajadorId":
            proveedor_trabajador_id,

        "rutaLecturaId":
            fila_ruta_id,

        "sectorId":
            fila_sector_id,

        "serieMedidor":
            None,

        "unidadEmpresarialId":
            unidad_id
    }

    headers = {

        "Authorization":
            f"Bearer {token}",

        "Content-Type":
            "application/json; charset=utf-8",

        "x-audit-application":
            "WEB",

        "x-audit-username":
            "72690389"
    }

    try:

        respuesta = requests.post(
            URL_MAPA_RUTA,
            json=payload,
            headers=headers,
            timeout=60
        )

        if respuesta.status_code != 200:
            return []

        datos = respuesta.json()

        selfies = datos.get(
            "selfies",
            []
        )

        resultado = []

        for selfie in selfies:

            selfie_url = selfie.get(
                "selfieUrl"
            )

            if not selfie_url:
                continue

            resultado.append({

                "lecturista":
                    selfie.get(
                        "lecturista",
                        ""
                    ),

                "ordenTrabajoAsignacionSelfieId":
                    selfie.get(
                        "ordenTrabajoAsignacionSelfieId"
                    ),

                "ordenTrabajoAsignacionId":
                    selfie.get(
                        "ordenTrabajoAsignacionId"
                    ),

                "ordenTrabajoAsignacionDetalleId":
                    selfie.get(
                        "ordenTrabajoAsignacionDetalleId"
                    ),

                "proveedorTrabajadorId":
                    selfie.get(
                        "proveedorTrabajadorId"
                    ),

                "secuenciaEnRuta":
                    selfie.get(
                        "secuenciaEnRuta"
                    ),

                "totalRuta":
                    selfie.get(
                        "totalRuta"
                    ),

                "fechaSelfie":
                    selfie.get(
                        "fechaSelfie"
                    ),

                "gpsx":
                    selfie.get(
                        "gpsx"
                    ),

                "gpsy":
                    selfie.get(
                        "gpsy"
                    ),

                "rutaLecturaId":
                    selfie.get(
                        "rutaLecturaId"
                    ),

                "rutaLecturaNombre":
                    selfie.get(
                        "rutaLecturaNombre",
                        ""
                    ),

                "selfieUrl":
                    selfie_url
            })

        return resultado

    except requests.exceptions.RequestException:
        return []

    except ValueError:
        return []


# ============================================================
# OBTENER SELFIES DE LAS FILAS DEL RESUMEN
# ============================================================

def obtener_selfies_de_filas(
    token,
    filas,
    periodo,
    proveedor_id,
    unidad_empresarial_id,
    ciclo_id,
    sector_id,
    ruta_lectura_id
):

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:
        return []

    todas_las_selfies = []

    consultas_realizadas = set()

    for fila in filas:

        asignaciones = fila.get(
            "asignaciones",
            []
        )

        if not asignaciones:
            continue

        for asignacion in asignaciones:

            asignacion_id = asignacion.get(
                "ordenTrabajoAsignacionId"
            )

            if not asignacion_id:
                continue

            proveedor_trabajador_id = asignacion.get(
                "proveedorTrabajadorId"
            )

            if not proveedor_trabajador_id:

                proveedor_trabajador_id = fila.get(
                    "proveedorTrabajadorId"
                )

            if not proveedor_trabajador_id:

                proveedor_trabajador_id = fila.get(
                    "proveedorTrabajadorID"
                )

            fecha_lectura = obtener_valor(
                asignacion,
                [
                    "fechaLectura",
                    "fecha",
                    "fechaAsignacion"
                ]
            )

            if fecha_lectura in (
                None,
                ""
            ):

                fecha_lectura = obtener_valor(
                    fila,
                    [
                        "fechaLectura",
                        "fecha",
                        "fechaAsignacion"
                    ]
                )

            if proveedor_trabajador_id in (
                None,
                ""
            ):
                continue

            if asignacion_id in (
                None,
                ""
            ):
                continue

            if fecha_lectura in (
                None,
                ""
            ):
                continue

            fecha_api = fecha_para_api(
                fecha_lectura
            )

            if not fecha_api:
                continue

            clave_consulta = (
                proveedor_trabajador_id,
                asignacion_id,
                fecha_api
            )

            if clave_consulta in consultas_realizadas:
                continue

            consultas_realizadas.add(
                clave_consulta
            )

            fila_asignacion = {

                "proveedorTrabajadorId":
                    proveedor_trabajador_id,

                "ordenTrabajoAsignacionId":
                    asignacion_id,

                "fechaLectura":
                    fecha_lectura,

                "fechaAsignacion":
                    fecha_lectura,

                "rutaLecturaId":
                    asignacion.get(
                        "rutaLecturaId"
                    ),

                "sectorId":
                    asignacion.get(
                        "sectorId"
                    )
            }

            selfies = obtener_selfies_asignacion(

                token=token,

                periodo=periodo,

                unidad_id=unidad_empresarial_id,

                fila=fila_asignacion,

                ciclo_id=ciclo_id,

                sector_id=sector_id,

                ruta_id=ruta_lectura_id
            )

            todas_las_selfies.extend(
                selfies
            )

    # ========================================================
    # ELIMINAR URLs REPETIDAS
    # ========================================================

    resultado_url = []

    urls = set()

    for selfie in todas_las_selfies:

        url = selfie.get(
            "selfieUrl"
        )

        if not url:
            continue

        if url in urls:
            continue

        urls.add(
            url
        )

        resultado_url.append(
            selfie
        )

    # ========================================================
    # ELIMINAR SECUENCIAS REPETIDAS
    # ========================================================

    resultado = []

    secuencias_vistas = set()

    for selfie in resultado_url:

        lecturista = selfie.get(
            "lecturista",
            ""
        )

        fecha_selfie = selfie.get(
            "fechaSelfie",
            ""
        )

        secuencia = selfie.get(
            "secuenciaEnRuta"
        )

        total_ruta = selfie.get(
            "totalRuta"
        )

        clave_secuencia = (

            str(lecturista),

            str(fecha_selfie),

            str(secuencia),

            str(total_ruta)
        )

        if clave_secuencia in secuencias_vistas:

            continue

        secuencias_vistas.add(
            clave_secuencia
        )

        resultado.append(
            selfie
        )

    return resultado


# ============================================================
# MOSTRAR SELFIES COMO GALERÍA
# AGRUPADA POR LECTURISTA Y DÍA
# ============================================================

def mostrar_galeria_selfies(selfies):

    if not selfies:

        st.info(
            "No se encontraron selfies para "
            "los registros seleccionados."
        )

        return

    grupos = {}

    for selfie in selfies:

        fecha_selfie = selfie.get(
            "fechaSelfie",
            ""
        )

        if not fecha_selfie:

            fecha = "Fecha desconocida"

        else:

            try:

                fecha = pd.to_datetime(
                    fecha_selfie
                ).strftime("%d/%m/%Y")

            except Exception:

                fecha = str(
                    fecha_selfie
                )[:10]

        lecturista = selfie.get(
            "lecturista",
            "Lecturista desconocido"
        )

        clave = (
            fecha,
            lecturista
        )

        if clave not in grupos:

            grupos[clave] = []

        grupos[clave].append(
            selfie
        )

    for (
        (fecha, lecturista),
        fotos_lecturista
    ) in grupos.items():

        col_info, col_fotos = st.columns(
            [1.2, 4.8]
        )

        with col_info:

            st.markdown(
                f"**{lecturista}**"
            )

            st.caption(
                f"📅 {fecha}"
            )

            st.caption(
                f"{len(fotos_lecturista)} selfie"
                +
                (
                    "s"
                    if len(fotos_lecturista) != 1
                    else ""
                )
            )

        with col_fotos:

            columnas = st.columns(6)

            for numero, selfie in enumerate(
                fotos_lecturista
            ):

                with columnas[
                    numero % 6
                ]:

                    st.image(
                        selfie["selfieUrl"],
                        use_container_width=True
                    )

                    # ----------------------------------------
                    # OBSERVACIÓN DE LA SELFIE
                    # ----------------------------------------

                    key_obs = clave_observacion(
                        selfie
                    )

                    st.markdown(
                        "<div style='text-align:center;'>"
                        "Observación"
                        "</div>",
                        unsafe_allow_html=True
                    )

                    st.selectbox(
                        "Observación",
                        OBSERVACIONES_LISTA,
                        index=0,
                        key=key_obs,
                        label_visibility="collapsed"
                    )

        st.divider()


# ============================================================
# CONVERTIR NÚMERO DE COLUMNA A LETRA DE EXCEL
# ============================================================

def letra_columna(numero):

    resultado = ""

    while numero > 0:

        numero, residuo = divmod(
            numero - 1,
            26
        )

        resultado = (
            chr(65 + residuo)
            + resultado
        )

    return resultado


# ============================================================
# GENERAR EXCEL DE SELFIES
# ============================================================

def generar_excel_selfies(selfies):

    if not selfies:
        return None

    grupos = {}

    for selfie in selfies:

        fecha_selfie = selfie.get(
            "fechaSelfie",
            ""
        )

        if not fecha_selfie:

            fecha = "Fecha desconocida"

        else:

            try:

                fecha = pd.to_datetime(
                    fecha_selfie
                ).strftime("%d/%m/%Y")

            except Exception:

                fecha = str(
                    fecha_selfie
                )[:10]

        lecturista = selfie.get(
            "lecturista",
            "Lecturista desconocido"
        )

        clave = (
            fecha,
            lecturista
        )

        if clave not in grupos:

            grupos[clave] = []

        grupos[clave].append(
            selfie
        )

    max_imgs = max(
        len(fotos)
        for fotos in grupos.values()
    )

    output = BytesIO()

    wb = Workbook()

    ws = wb.active

    ws.title = "Selfies"

    # ========================================================
    # ENCABEZADOS
    # ========================================================

    headers = [
        "Fecha Selfie",
        "Lecturista"
    ]

    for i in range(1, max_imgs + 1):

        headers.append(
            f"Url_foto_{i}"
        )

    for i in range(1, max_imgs + 1):

        headers.append(
            f"Imagen {i:02d}"
        )

    headers.append(
        "Observaciones"
    )

    ws.append(
        headers
    )

    # ========================================================
    # DATOS
    # ========================================================

    fila_inicio = 2
    primera_columna_imagen = 3

    for fila_idx, (
        (fecha, lecturista),
        fotos_lecturista
    ) in enumerate(
        grupos.items(),
        start=2
    ):

        fila = [
            fecha,
            lecturista
        ]

        # ----------------------------------------------------
        # URL DE LAS FOTOS
        # ----------------------------------------------------

        for i in range(max_imgs):

            if i < len(fotos_lecturista):

                url = fotos_lecturista[i].get(
                    "selfieUrl",
                    ""
                )

            else:

                url = ""

            fila.append(
                url
            )

        primera_columna_imagen = 3

        for i in range(max_imgs):

            col_letra = letra_columna(
                primera_columna_imagen + i
            )

            if fila_idx == fila_inicio and i == 0:

                fila.append(
                    f"'=SI.ERROR(IMAGEN({col_letra}{fila_inicio};;3;250;180);\"\")"
                )

            else:

                fila.append("")

        # ----------------------------------------------------
        # OBSERVACIONES
        # ----------------------------------------------------

        observaciones = {}

        for numero, selfie in enumerate(
            fotos_lecturista,
            start=1
        ):

            key_obs = clave_observacion(
                selfie
            )

            observacion = st.session_state.get(
                key_obs,
                "CORRECTO"
            )

            if observacion != "CORRECTO":

                if observacion not in observaciones:

                    observaciones[
                        observacion
                    ] = []

                observaciones[
                    observacion
                ].append(
                    numero
                )

        if not observaciones:

            texto_observaciones = "CORRECTO"

        else:

            partes = []

            for (
                observacion,
                numeros
            ) in observaciones.items():

                if len(numeros) == 1:

                    imagenes = (
                        f"imagen {numeros[0]:02d}"
                    )

                elif len(numeros) == 2:

                    imagenes = (
                        f"imagen {numeros[0]:02d} "
                        f"y {numeros[1]:02d}"
                    )

                else:

                    numeros_texto = ", ".join(
                        f"{n:02d}"
                        for n in numeros[:-1]
                    )

                    imagenes = (
                        f"imagen {numeros_texto} "
                        f"y {numeros[-1]:02d}"
                    )

                partes.append(
                    f"{imagenes}: {observacion}"
                )

            texto_observaciones = "\n".join(
                partes
            )

        fila.append(
            texto_observaciones
        )

        ws.append(
            fila
        )

    # ========================================================
    # FORMATO
    # ========================================================

    # URL ocultas

    primera_columna_url = 3

    ultima_columna_url = (
        2 + max_imgs
    )

    for columna in range(
        primera_columna_url,
        ultima_columna_url + 1
    ):

        letra = letra_columna(
            columna
        )

        ws.column_dimensions[
            letra
        ].hidden = True

    # Ancho de columnas

    ws.column_dimensions[
        "A"
    ].width = 18

    ws.column_dimensions[
        "B"
    ].width = 25

    primera_columna_imagen = (
        3 + max_imgs
    )

    for i in range(max_imgs):

        columna = (
            primera_columna_imagen
            + i
        )

        letra = letra_columna(
            columna
        )

        ws.column_dimensions[
            letra
        ].width = 25

    columna_observaciones = (
        primera_columna_imagen
        + max_imgs
    )

    letra_observaciones = letra_columna(
        columna_observaciones
    )

    ws.column_dimensions[
        letra_observaciones
    ].width = 50

    # Altura de filas

    for fila in range(
        2,
        ws.max_row + 1
    ):

        ws.row_dimensions[
            fila
        ].height = 189
    
    # Ajuste de observaciones

    for fila in range(
        2,
        ws.max_row + 1
    ):

        ws.cell(
            fila,
            columna_observaciones
        ).alignment = (
            ws.cell(
                fila,
                columna_observaciones
            ).alignment.copy(
                wrap_text=True,
                vertical="top"
            )
        )

    wb.save(
        output
    )

    output.seek(0)

    return output.getvalue()


# ============================================================
# MOSTRAR RESULTADOS
# ============================================================

def mostrar_resultados_selfies(selfies):

    st.subheader(
        "📸 Selfies"
    )

    st.write(
        f"Cantidad obtenida: {len(selfies)}"
    )

    mostrar_galeria_selfies(
        selfies
    )

    if selfies:

        archivo_excel = generar_excel_selfies(
            selfies
        )

        st.download_button(

            "📥 Descargar Selfis",

            data=archivo_excel,

            file_name="Reporte_Selfie.xlsx",

            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),

            key="descargar_reporte_selfies"
        )


# ============================================================
# CONSULTAR LECTURAS
# ============================================================

def consultar_lecturas():

    token = st.session_state.get(
        "fieldservice_token"
    )

    if not token:

        st.error(
            "No existe una sesión activa de FieldService."
        )

        return

    periodo = st.session_state.get(
        "fs_periodo"
    )

    unidad_id = st.session_state.get(
        "fs_unidad"
    )

    fecha_desde = st.session_state.get(
        "fs_fecha_desde"
    )

    fecha_hasta = st.session_state.get(
        "fs_fecha_hasta"
    )

    ciclo_id = st.session_state.get(
        "fs_ciclo_id"
    )

    sector_id = st.session_state.get(
        "fs_sector_id"
    )

    ruta_id = st.session_state.get(
        "fs_ruta_id"
    )

    lecturista_id = st.session_state.get(
        "fs_lecturista_id"
    )

    if unidad_id is None:

        st.warning(
            "Seleccione una unidad de negocio."
        )

        return

    if periodo is None:

        st.warning(
            "Seleccione un periodo."
        )

        return

    if fecha_desde is None:
        fecha_desde = date.today()

    if fecha_hasta is None:
        fecha_hasta = date.today()

    fecha_desde_api = (
        f"{fecha_desde.isoformat()}T00:00:00"
    )

    fecha_hasta_api = (
        f"{fecha_hasta.isoformat()}T00:00:00"
    )

    payload = {

        "cicloId":
            ciclo_id,

        "fechaAsignacion":
            None,

        "fechaLecturaDesde":
            fecha_desde_api,

        "fechaLecturaHasta":
            fecha_hasta_api,

        "nroServicioDigitado":
            None,

        "ordenTrabajoAsignacionIds":
            None,

        "ordenTrabajoId":
            None,

        "periodo":
            periodo,

        "proveedorId":
            3,

        "proveedorTrabajadorId":
            lecturista_id,

        "rutaLecturaId":
            ruta_id,

        "sectorId":
            sector_id,

        "serieMedidor":
            None,

        "unidadEmpresarialId":
            unidad_id
    }

    headers = {

        "Authorization":
            f"Bearer {token}",

        "Content-Type":
            "application/json; charset=utf-8",

        "x-audit-application":
            "WEB",

        "x-audit-username":
            "72690389"
    }

    try:

        respuesta = requests.post(
            URL_CONSULTAR,
            json=payload,
            headers=headers,
            timeout=60
        )

        if respuesta.status_code != 200:

            st.error(
                "Error en la consulta. "
                f"Código HTTP: {respuesta.status_code}"
            )

            st.write(
                "Respuesta del servidor:",
                respuesta.text
            )

            return None

        datos = respuesta.json()

        st.session_state.fs_resultado = datos

        st.success(
            "Consulta realizada correctamente"
        )

        # ----------------------------------------------------
        # OBTENER LECTURISTAS
        # ----------------------------------------------------

        lecturistas_resultado = datos.get(
            "lecturistas",
            []
        )

        if lecturista_id is not None:

            filas_para_selfies = [

                item

                for item in lecturistas_resultado

                if (
                    item.get(
                        "proveedorTrabajadorId"
                    )
                    == lecturista_id

                    or

                    item.get(
                        "proveedorTrabajadorID"
                    )
                    == lecturista_id

                    or

                    item.get(
                        "lecturistaId"
                    )
                    == lecturista_id

                    or

                    item.get(
                        "trabajadorId"
                    )
                    == lecturista_id

                    or

                    item.get(
                        "id"
                    )
                    == lecturista_id
                )
            ]

        else:

            filas_para_selfies = (
                lecturistas_resultado
            )

        # ----------------------------------------------------
        # CONSULTAR SELFIES
        # ----------------------------------------------------

        selfies = obtener_selfies_de_filas(

            token=token,

            filas=filas_para_selfies,

            periodo=periodo,

            proveedor_id=3,

            unidad_empresarial_id=unidad_id,

            ciclo_id=ciclo_id,

            sector_id=sector_id,

            ruta_lectura_id=ruta_id
        )

        st.session_state.fs_selfies = selfies

        # ----------------------------------------------------
        # MOSTRAR RESULTADOS
        # ----------------------------------------------------

        mostrar_resultados_selfies(
            selfies
        )

        return datos

    except requests.exceptions.RequestException as e:

        st.error(
            f"Error de conexión con FieldService: {e}"
        )

        return None

    except ValueError:

        st.error(
            "La respuesta de FieldService no tiene "
            "un formato JSON válido."
        )

        return None


# ============================================================
# EJECUTAR FIELDSERVICE
# ============================================================

def ejecutar_fieldservice():

    # ========================================================
    # INICIALIZAR SESSION STATE
    # ========================================================

    if "fieldservice_token" not in st.session_state:

        st.session_state.fieldservice_token = None

    if "fs_resultado" not in st.session_state:

        st.session_state.fs_resultado = None

    if "fs_selfies" not in st.session_state:

        st.session_state.fs_selfies = []

    # ========================================================
    # LOGIN
    # ========================================================

    if not st.session_state.fieldservice_token:

        st.subheader(
            "🔐 Acceso a FieldService"
        )

        usuario = st.text_input(
            "Usuario",
            key="fieldservice_usuario"
        )

        clave = st.text_input(
            "Contraseña",
            type="password",
            key="fieldservice_clave"
        )

        if st.button(
            "Iniciar sesión",
            type="primary"
        ):

            if not usuario or not clave:

                st.warning(
                    "Ingrese usuario y contraseña."
                )

            else:

                token = login_fieldservice(
                    usuario,
                    clave
                )

                if token:

                    st.session_state.fieldservice_token = (
                        token
                    )

                    st.success(
                        "🟢 Inicio de sesión correcto."
                    )

                    st.rerun()

                else:

                    st.error(
                        "❌ Usuario o contraseña incorrectos, "
                        "o no fue posible conectarse con FieldService."
                    )

        return

    # ========================================================
    # SESIÓN ACTIVA
    # ========================================================

    st.subheader(
        "📊 Ejecución y Monitoreo"
    )

    # ========================================================
    # FILA 1
    # ========================================================

    col_unidad, col_periodo, col_fecha_desde = st.columns(3)

    with col_unidad:

        unidades = {

            76: "Ayacucho",

            77: "Huancayo",

            78: "Huancavelica",

            79: "Tarma",

            80: "Selva Central",

            81: "Pasco",

            82: "Huánuco",

            83: "Valle Mantaro",

            84: "Tingo María"
        }

        unidad_id = st.selectbox(

            "Unidad de negocio",

            options=list(
                unidades.keys()
            ),

            index=None,

            placeholder="Seleccione",

            format_func=lambda x:
                unidades[x],

            key="fs_unidad"
        )

    with col_periodo:

        periodos = [
            202608
        ]

        periodo = st.selectbox(

            "Periodo",

            options=periodos,

            index=None,

            placeholder="Seleccione",

            key="fs_periodo"
        )

    with col_fecha_desde:

        fecha_desde = st.date_input(

            "Fecha lectura desde",

            key="fs_fecha_desde"
        )

    # ========================================================
    # FILA 2
    # ========================================================

    col_fecha_hasta, col_ciclo, col_sector = st.columns(3)

    with col_fecha_hasta:

        fecha_hasta = st.date_input(

            "Fecha lectura hasta",

            value=date.today(),

            key="fs_fecha_hasta"
        )

    with col_ciclo:

        ciclos = []

        if unidad_id is not None:

            ciclos = obtener_ciclos(
                unidad_id
            )

        opciones_ciclos = {

            None: "Todos",

            **{

                item["id"]: item["nombre"]

                for item in ciclos

                if item.get("id") is not None

                and item.get("nombre")
            }
        }

        ciclo_id = st.selectbox(

            "Ciclo",

            options=list(
                opciones_ciclos.keys()
            ),

            index=0,

            placeholder="Buscar ciclo por id o nombre",

            format_func=lambda x:
                opciones_ciclos[x],

            key="fs_ciclo"
        )

        st.session_state.fs_ciclo_id = (
            ciclo_id
        )

    with col_sector:

        sectores = {
            None: "Todos"
        }

        if (
            unidad_id is not None
            and ciclo_id is not None
        ):

            datos_sectores = obtener_sectores(
                unidad_id,
                ciclo_id
            )

            for item in datos_sectores:

                sector_id_item = item.get(
                    "id"
                )

                sector_nombre_item = item.get(
                    "nombre"
                )

                if (
                    sector_id_item is not None
                    and sector_nombre_item
                ):

                    sectores[sector_id_item] = (
                        sector_nombre_item
                    )

        sector_id = st.selectbox(

            "Sector",

            options=list(
                sectores.keys()
            ),

            index=0,

            format_func=lambda x:
                sectores[x],

            key="fs_sector"
        )

        st.session_state.fs_sector_id = (
            sector_id
        )

    # ========================================================
    # FILA 3
    # ========================================================

    col_ruta, col_lecturista, col_buscar = st.columns(3)

    with col_ruta:

        rutas = {
            None: "Todas"
        }

        if (
            unidad_id is not None
            and sector_id is not None
        ):

            datos_rutas = obtener_rutas(
                unidad_id,
                sector_id
            )

            for item in datos_rutas:

                ruta_id_item = item.get(
                    "id"
                )

                ruta_nombre_item = item.get(
                    "nombre"
                )

                if (
                    ruta_id_item is not None
                    and ruta_nombre_item
                ):

                    rutas[ruta_id_item] = (
                        f"{ruta_id_item} - "
                        f"{ruta_nombre_item}"
                    )

        ruta_id = st.selectbox(

            "Ruta de lectura",

            options=list(
                rutas.keys()
            ),

            index=0,

            format_func=lambda x:
                rutas[x],

            key="fs_ruta"
        )

        st.session_state.fs_ruta_id = (
            ruta_id
        )

    with col_lecturista:

        lecturistas = []

        if (
            unidad_id is not None
            and periodo is not None
        ):

            lecturistas = obtener_lecturistas(
                unidad_id,
                periodo
            )

        opciones_lecturistas = {

            item["id"]: item["nombre"]

            for item in lecturistas

            if item.get("id") is not None

            and item.get("nombre")
        }

        opciones_lecturistas_con_todos = {

            None: "Todos",

            **opciones_lecturistas
        }

        lecturista_id = st.selectbox(

            "Lecturista",

            options=list(
                opciones_lecturistas_con_todos.keys()
            ),

            index=0,

            format_func=lambda x:
                opciones_lecturistas_con_todos[x],

            key="fs_lecturista"
        )

        st.session_state.fs_lecturista_id = (
            lecturista_id
        )

    with col_buscar:

        st.write("")
        st.write("")

        buscar = st.button(

            "🔎 Buscar",

            type="primary",

            use_container_width=True
        )

    # ========================================================
    # BUSCAR
    # ========================================================

    if buscar:

        if unidad_id is None:

            st.warning(
                "Seleccione una unidad de negocio."
            )

        elif periodo is None:

            st.warning(
                "Seleccione un periodo."
            )

        elif fecha_desde > fecha_hasta:

            st.warning(
                "La fecha desde no puede ser mayor "
                "que la fecha hasta."
            )

        else:

            consultar_lecturas()

    # ========================================================
    # MANTENER RESULTADOS AL CAMBIAR OBSERVACIONES
    # ========================================================

    elif st.session_state.fs_resultado is not None:

        mostrar_resultados_selfies(
            st.session_state.fs_selfies
        )